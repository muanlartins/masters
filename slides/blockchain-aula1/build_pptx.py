"""Build apresentacao_blockchain_aula1.pptx — Aula 1 (~2 h, portugues).

Por que blockchain existe: do hash ao consenso. Segue PRESENTATION_GUIDE §1:
cada slide de conteudo = um titulo-frase (topo, esquerda, declarativo) + um
artefato visual + espaco em branco. Sem bullets, sem caixas de texto, sem
rodape/logo. O professor narra tudo que o slide nao mostra. Plano: PLAN.md.

Slides de demo ao vivo (Anders) e o roleplay carregam um diagrama-ancora de
backup; a aula em si acontece no navegador / na atividade.

Estilo herdado de slides/eniac-waldo/build_pptx.py.
Run: venv/bin/python slides/blockchain/build_pptx.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path
from PIL import Image

HERE = Path(__file__).parent
ASSETS = HERE / "assets"
OUTPUT = HERE / "apresentacao_blockchain_aula1.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN_X = Inches(0.6)

BLACK = RGBColor(0x1f, 0x29, 0x37)
GRAY = RGBColor(0x6b, 0x72, 0x80)
WHITE = RGBColor(0xff, 0xff, 0xff)
CHAIN = RGBColor(0x1f, 0x77, 0xb4)
GREEN = RGBColor(0x2c, 0xa0, 0x2c)
RED = RGBColor(0xd6, 0x27, 0x28)
ENERGY = RGBColor(0xff, 0x7f, 0x0e)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid(); bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background(); bg.shadow.inherit = False
    return s


def title(s, text, size=30, color=BLACK):
    tb = s.shapes.add_textbox(MARGIN_X, Inches(0.34),
                              SLIDE_W - 2 * MARGIN_X, Inches(1.1))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0); tf.margin_top = Inches(0)
    p = tf.paragraphs[0]; p.text = text; p.alignment = PP_ALIGN.LEFT
    r = p.runs[0]; r.font.size = Pt(size); r.font.bold = True
    r.font.name = "Helvetica"; r.font.color.rgb = color


def centered(s, lines, top, height, sizes, bolds=None, colors=None):
    n = len(lines)
    bolds = bolds or [False] * n
    colors = colors or [BLACK] * n
    tb = s.shapes.add_textbox(MARGIN_X, top, SLIDE_W - 2 * MARGIN_X, height)
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0); tf.margin_top = Inches(0)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln; p.alignment = PP_ALIGN.CENTER
        if i > 0:
            p.space_before = Pt(14)
        r = p.runs[0]
        r.font.size = Pt(sizes[i]); r.font.bold = bolds[i]
        r.font.name = "Helvetica"; r.font.color.rgb = colors[i]


def image(s, name, top=Inches(1.75), max_h=Inches(5.1), max_w=None):
    path = ASSETS / name
    if max_w is None:
        max_w = SLIDE_W - 2 * MARGIN_X
    iw, ih = Image.open(path).size; ratio = iw / ih
    h = max_h; w = Emu(int(h * ratio))
    if w > max_w:
        w = max_w; h = Emu(int(w / ratio))
    left = Emu(int((SLIDE_W - w) / 2))
    s.shapes.add_picture(str(path), left, top, w, h)


def divider(s, text):
    centered(s, [text], Inches(3.0), Inches(1.5), [34], bolds=[True])


def subtitle(s, text, top=6.55, size=15, color=GRAY):
    """Legenda editável: texto de slide (mexa no Keynote), não embutido na figura."""
    tb = s.shapes.add_textbox(MARGIN_X, Inches(top), SLIDE_W - 2 * MARGIN_X, Inches(0.55))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text; p.alignment = PP_ALIGN.CENTER
    r = p.runs[0]; r.font.size = Pt(size); r.font.name = "Helvetica"
    r.font.italic = True; r.font.color.rgb = color


def demo_badge(s, text="▶ demo ao vivo"):
    """Lembrete sutil (chip azul, canto sup. dir.): este slide tem demo ao vivo."""
    w = Inches(0.55 + 0.097 * len(text)); h = Inches(0.42)
    left = SLIDE_W - MARGIN_X - w
    chip = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(0.4), w, h)
    chip.fill.solid(); chip.fill.fore_color.rgb = RGBColor(0xe8, 0xf1, 0xfa)
    chip.line.color.rgb = CHAIN; chip.line.width = Pt(1.25)
    chip.shadow.inherit = False
    tf = chip.text_frame; tf.word_wrap = False
    tf.margin_top = Inches(0.02); tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]; p.text = text; p.alignment = PP_ALIGN.CENTER
    r = p.runs[0]; r.font.size = Pt(13); r.font.bold = True
    r.font.name = "Helvetica"; r.font.color.rgb = CHAIN


def content(title_text, img, size=28, top=1.75, max_h=5.1, sub=None, demo=None):
    s = slide(); title(s, title_text, size=size)
    image(s, img, top=Inches(top), max_h=Inches(max_h))
    if sub:
        subtitle(s, sub)
    if demo:
        demo_badge(s, demo)
    return s


# =========================================================== 1  capa
s = slide()
centered(s, ["Blockchain, do zero:", "por que ela existe"],
         Inches(2.2), Inches(2.0), [40, 40], bolds=[True, True])
centered(s, ["Programação Avançada  ·  6º período",
             "Aula 1 de 2",
             "do hash ao consenso"],
         Inches(4.5), Inches(2.2), [22, 20, 18],
         colors=[BLACK, GRAY, CHAIN])

# motivacao
content("Dinheiro digital é só um arquivo — e arquivo se copia",
        "diagram_double_spend.png", size=26, top=1.9, max_h=4.3,
        sub="a mesma moeda paga Bob e Carol — o problema do gasto-duplo")
s = content("Um banco resolve isso — e te devolve um dono no meio",
            "diagram_central_ledger.png", size=26, top=1.7, max_h=4.0)
subtitle(s, "um banco de dados comum (um dono) resolve o gasto-duplo na hora",
         top=5.95, color=GREEN)
subtitle(s, "mas vira ponto único de confiança · censura · falha",
         top=6.5, color=RED)
s = content("Tira o banco: todo mundo guarda uma cópia do livro",
            "diagram_distributed_ledger.png", size=27, top=1.55, max_h=4.2)
subtitle(s, "todos guardam uma cópia do mesmo livro", top=5.95, color=GRAY)
subtitle(s, "mas como milhares de cópias concordam sobre o que entrou?",
         top=6.5, color=RED)

# divisor cripto
divider(slide(), "Antes, duas peças de criptografia")
content("Um hash é uma impressão digital de mão única",
        "diagram_hash.png", size=28, demo="▶ demo · Anders /hash")
s = content("Uma assinatura digital prova quem autorizou",
            "diagram_signature.png", size=28, top=1.7, max_h=4.1,
            demo="▶ demo · Anders /signatures")
subtitle(s, "dá pra conferir sem poder forjar — é assim que se autoriza uma transação",
         top=6.25, color=CHAIN)

# divisor cadeia
divider(slide(), "Do hash à cadeia")
s = content("Encadear pelo hash do anterior torna a fraude evidente",
            "diagram_hashchain.png", size=26, top=1.7, max_h=4.0,
            demo="▶ demo · Anders /blockchain")
subtitle(s, "cada bloco guarda o hash do anterior", top=5.95, color=GRAY)
subtitle(s, "mexeu em um → o elo não bate e todos os seguintes quebram",
         top=6.5, color=RED)
s = content("Minerar é achar um número que faça o hash começar com zeros",
            "diagram_mining.png", size=24, top=1.65, max_h=4.0,
            demo="▶ demo · Anders /block + Mine")
subtitle(s, "sem atalho: testa nonce após nonce até o hash começar com 0000",
         top=5.95, color=GRAY)
subtitle(s, "esse trabalho é o que vai custar energia", top=6.5, color=ENERGY)

# divisor consenso
divider(slide(), "Quem escolhe o próximo bloco?")
# generais bizantinos: funde o problema (12) + a conclusão 3f+1 (13) num slide só
s = content("Num bando de desconhecidos, em quem você confia?",
            "diagram_generals.png", size=26, top=1.7, max_h=4.0,
            demo="▶ roleplay · Generais")
subtitle(s, "o traidor conta versões diferentes para cada general",
         top=5.95, color=GRAY)
subtitle(s, "ninguém sabe quem mente → só resolve com mais de 2/3 honestos (3f+1)",
         top=6.5, color=RED)

# intervalo — logo após o roleplay, quebra no cliffhanger do consenso
s = slide()
centered(s, ["Intervalo"], Inches(2.5), Inches(1.4), [46], bolds=[True])
centered(s, ["~10 min"], Inches(3.95), Inches(0.8), [24], colors=[GRAY])
centered(s, ["na volta: a resposta de 2008 — como o Bitcoin resolve o impossível"],
         Inches(5.05), Inches(0.8), [18], colors=[CHAIN])

# divisor resposta — a sacada conceitual
divider(slide(), "A resposta de 2008")
content('Proof of work: tornar "votar" caro a ponto de não dar pra fingir',
        "diagram_sybil_pow.png", size=24)

# divisor mecanismo — abrir o capô (bloco técnico, pós-intervalo)
divider(slide(), "Abrindo o capô: as peças, por dentro")

# 1. árvore de Merkle
s = content("A árvore de Merkle prova que uma transação está no bloco",
            "fig_merkle.png", size=26, top=1.75, max_h=4.2)
subtitle(s, "…sem baixá-lo inteiro: a carteira recompõe a raiz com só log(n) hashes, não as milhares de txs",
         top=6.25, color=CHAIN)
# 2. cabeçalho de 80 bytes + o elo
s = content("O cabeçalho de 80 bytes é o que se faz hash — e é o que encadeia",
            "fig_header.png", size=24, top=1.6, max_h=4.1)
subtitle(s, "o hash do bloco é o SHA256 (em dobro) do cabeçalho — e vira o “prev” do próximo",
         top=6.15, color=GRAY)
# 3. prova de trabalho: hash < alvo
s = content("Minerar é achar um cabeçalho cujo hash caia abaixo de um alvo",
            "fig_pow_target.png", size=24, top=1.7, max_h=3.9,
            demo="▶ retoma Anders /block")
subtitle(s, "é o “0000” do simulador, agora formal — e sem atalho a não ser tentar",
         top=6.05, color=ENERGY)
# 4. dificuldade: retarget
s = content("A dificuldade se reajusta para manter ~10 minutos por bloco",
            "fig_difficulty.png", size=25, top=1.7, max_h=4.0)
subtitle(s, "mais poder de mineração na rede → alvo desce → o ritmo volta a ~10 min",
         top=6.1, color=CHAIN)
# 5. cadeia mais longa (mantém a demo Anders /distributed)
s = content("Vence a cadeia com mais trabalho acumulado",
            "diagram_longest_chain.png", size=27, top=1.7, max_h=4.0,
            demo="▶ demo · Anders /distributed")
subtitle(s, "sem autoridade, o desempate é objetivo: a história com mais prova de trabalho",
         top=6.1, color=GRAY)
# 6. segurança: reverter custa o trabalho de volta
s = content("Reescrever a história significa refazer todo o trabalho",
            "fig_security.png", size=25, top=1.7, max_h=4.0)
subtitle(s, "atacante com menos da metade do poder: a chance de reverter cai exponencialmente com a profundidade",
         top=6.1, color=RED)

# síntese: inventário das peças — todas já existiam, menos uma
s = content("O Bitcoin não inventou as peças — inventou juntá-las",
            "fig_timeline.png", size=26, top=1.55, max_h=4.7)
subtitle(s, "tudo à esquerda já existia (até a visão de moeda PoW); faltava fazê-las concordar numa rede aberta, sem saber quem participa",
         top=6.4, color=CHAIN)

# divisor custo
divider(slide(), "Quanto custa esse consenso?")
content("Esse consenso gasta a eletricidade de um país inteiro",
        "fig_energy.png", size=26, top=1.85, max_h=4.9)
content("E faz ~7 transações por segundo, não 65 mil",
        "fig_tps.png", size=28, top=1.85, max_h=4.9)

# final
content("Construímos um livro de moedas. E se ele guardasse código?",
        "diagram_bridge.png", size=25)

prs.save(OUTPUT)
print(f"wrote {OUTPUT}  ({len(prs.slides._sldIdLst)} slides)")
