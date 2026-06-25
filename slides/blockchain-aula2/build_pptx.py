"""Build apresentacao_blockchain_aula2.pptx — Aula 2 (~2 h, portugues).

Da moeda ao computador-mundo: smart contracts, gas, reentrancia/fork, Etherscan
e a pergunta de fecho ("precisa mesmo de blockchain?"). Segue PRESENTATION_GUIDE
§1: cada slide de conteudo = um titulo-frase (topo, esquerda, declarativo) + um
artefato visual + espaco em branco. Sem bullets, sem caixas de texto, sem
rodape/logo. O professor narra tudo que o slide nao mostra. Plano: PLAN.md.

Slides de demo ao vivo (Remix, Etherscan) carregam um diagrama-ancora de backup;
a aula em si acontece no navegador. Estilo herdado de slides/blockchain/build_pptx.py.
Run: venv/bin/python slides/blockchain-aula2/build_pptx.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path
from PIL import Image

HERE = Path(__file__).parent
ASSETS = HERE / "assets"
OUTPUT = HERE / "apresentacao_blockchain_aula2.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN_X = Inches(0.6)

BLACK = RGBColor(0x1f, 0x29, 0x37)
GRAY = RGBColor(0x6b, 0x72, 0x80)
WHITE = RGBColor(0xff, 0xff, 0xff)
CHAIN = RGBColor(0x1f, 0x77, 0xb4)
CODE = RGBColor(0x2c, 0xa0, 0x2c)
RED = RGBColor(0xd6, 0x27, 0x28)
ENERGY = RGBColor(0xff, 0x7f, 0x0e)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid(); bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background(); bg.shadow.inherit = False
    return s


def title(s, text, size=30, color=BLACK):
    tb = s.shapes.add_textbox(MARGIN_X, Inches(0.34),
                              SLIDE_W - 2 * MARGIN_X, Inches(1.1))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0); tf.margin_top = Inches(0)
    p = tf.paragraphs[0]; p.text = text; p.alignment = PP_ALIGN.LEFT
    r = p.runs[0]; r.font.size = Pt(size); r.font.bold = True
    r.font.name = "Helvetica"; r.font.color.rgb = color


def centered(s, lines, top, height, sizes, bolds=None, colors=None):
    n = len(lines)
    bolds = bolds or [False] * n
    colors = colors or [BLACK] * n
    tb = s.shapes.add_textbox(MARGIN_X, top, SLIDE_W - 2 * MARGIN_X, height)
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0); tf.margin_top = Inches(0)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln; p.alignment = PP_ALIGN.CENTER
        if i > 0:
            p.space_before = Pt(14)
        r = p.runs[0]
        r.font.size = Pt(sizes[i]); r.font.bold = bolds[i]
        r.font.name = "Helvetica"; r.font.color.rgb = colors[i]


def image(s, name, top=Inches(1.75), max_h=Inches(5.1), max_w=None):
    path = ASSETS / name
    if max_w is None:
        max_w = SLIDE_W - 2 * MARGIN_X
    iw, ih = Image.open(path).size; ratio = iw / ih
    h = max_h; w = Emu(int(h * ratio))
    if w > max_w:
        w = max_w; h = Emu(int(w / ratio))
    left = Emu(int((SLIDE_W - w) / 2))
    s.shapes.add_picture(str(path), left, top, w, h)


def divider(s, text):
    centered(s, [text], Inches(3.0), Inches(1.5), [34], bolds=[True])


def subtitle(s, text, top=6.55, size=15, color=GRAY):
    """Legenda editavel: texto de slide (mexa no Keynote), nao embutido na figura."""
    tb = s.shapes.add_textbox(MARGIN_X, Inches(top), SLIDE_W - 2 * MARGIN_X, Inches(0.55))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text; p.alignment = PP_ALIGN.CENTER
    r = p.runs[0]; r.font.size = Pt(size); r.font.name = "Helvetica"
    r.font.italic = True; r.font.color.rgb = color


def demo_badge(s, text="▶ demo ao vivo"):
    """Lembrete sutil (chip azul, canto sup. dir.): este slide tem demo ao vivo."""
    w = Inches(0.55 + 0.097 * len(text)); h = Inches(0.42)
    left = SLIDE_W - MARGIN_X - w
    chip = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(0.4), w, h)
    chip.fill.solid(); chip.fill.fore_color.rgb = RGBColor(0xe8, 0xf1, 0xfa)
    chip.line.color.rgb = CHAIN; chip.line.width = Pt(1.25)
    chip.shadow.inherit = False
    tf = chip.text_frame; tf.word_wrap = False
    tf.margin_top = Inches(0.02); tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]; p.text = text; p.alignment = PP_ALIGN.CENTER
    r = p.runs[0]; r.font.size = Pt(13); r.font.bold = True
    r.font.name = "Helvetica"; r.font.color.rgb = CHAIN


def content(title_text, img, size=28, top=1.75, max_h=5.1, sub=None, demo=None):
    s = slide(); title(s, title_text, size=size)
    image(s, img, top=Inches(top), max_h=Inches(max_h))
    if sub:
        subtitle(s, sub)
    if demo:
        demo_badge(s, demo)
    return s


# =========================================================== 1  capa
s = slide()
centered(s, ["Da moeda ao", "computador-mundo"],
         Inches(2.2), Inches(2.0), [40, 40], bolds=[True, True])
centered(s, ["Programação Avançada  ·  6º período",
             "Aula 2 de 2",
             "contratos inteligentes"],
         Inches(4.5), Inches(2.2), [22, 20, 18],
         colors=[BLACK, GRAY, CHAIN])

# ---- recap + a virada ----
divider(slide(), "Recap: o que a Aula 1 deixou pronto")
s = content("Na Aula 1, um livro de moedas; e se cada entrada guardasse código?",
            "diagram_recap.png", size=25, top=1.7, max_h=4.3)
subtitle(s, "o mesmo consenso da Aula 1 — mas agora o que se registra pode ser código",
         top=6.2, color=CODE)
s = content("Uma conta é um estado; uma transação é uma transição de estado",
            "diagram_state_machine.png", size=25, top=1.8, max_h=4.3)
s = content("Todo nó roda o mesmo código e concorda no resultado: um computador-mundo",
            "diagram_world_computer.png", size=23, top=1.75, max_h=4.3)
subtitle(s, "no Bitcoin a rede concorda no mesmo saldo; aqui ela concorda no mesmo resultado de cada execução",
         top=6.3, color=CHAIN)

# ---- contratos inteligentes ----
divider(slide(), "Contratos inteligentes")
s = content("Um contrato é código que mora na rede e se cumpre sozinho",
            "diagram_vending.png", size=26, top=1.6, max_h=4.7)
s = content("O que ele não consegue: enxergar qualquer coisa de fora da rede",
            "diagram_oracle.png", size=25, top=1.6, max_h=4.7)
s = content("Cada passo de código custa gas — é o que impede o loop infinito",
            "fig_gas.png", size=25, top=1.55, max_h=4.6)
subtitle(s, "sem preço, um while(true) travaria milhares de nós; é o gas que impõe um fim a toda execução",
         top=6.35, color=ENERGY)

# ---- mãos à obra: Remix ----
divider(slide(), "Mãos à obra: contratos rodando")
s = content("Counter: toda escrita é uma transação; toda leitura é de graça",
            "diagram_counter.png", size=25, top=1.7, max_h=4.0,
            demo="▶ demo · Remix VM")
subtitle(s, "increment = transação (gasta gas, vira bloco) · ler count = call (grátis) · decrement em 0 → revert",
         top=6.1, color=GRAY)
s = content("Escrow: confiança programável — só o árbitro libera o dinheiro",
            "diagram_escrow.png", size=25, top=1.7, max_h=4.1,
            demo="▶ demo · Remix VM")
subtitle(s, "payable/msg.value: o contrato guarda 1 ETH; nem comprador nem vendedor mexem — só o árbitro decide",
         top=6.25, color=CODE)

# ---- intervalo ----
s = slide()
centered(s, ["Intervalo"], Inches(2.5), Inches(1.4), [46], bolds=[True])
centered(s, ["~10 min"], Inches(3.95), Inches(0.8), [24], colors=[GRAY])
centered(s, ["na volta: o bug de 60 milhões de dólares — quando o contrato certo faz a coisa errada"],
         Inches(5.05), Inches(0.8), [18], colors=[RED])

# ---- o bug de 60 milhões: reentrância ----
divider(slide(), "O bug de 60 milhões de dólares")
s = content("Reentrância: o atacante volta a sacar antes de o saldo zerar",
            "diagram_reentrancy.png", size=25, top=1.7, max_h=4.1,
            demo="▶ demo · Remix VM")
subtitle(s, "o withdraw envia o ETH antes de zerar o saldo; o receive() do atacante reentra e saca de novo — depositou 1, levou 6",
         top=6.25, color=RED)
s = content("A correção é uma linha: zere o saldo antes de enviar",
            "diagram_checks_effects.png", size=26, top=1.75, max_h=4.0)
subtitle(s, "checks-effects-interactions: mude o estado antes de chamar para fora — a única diferença é a ordem",
         top=6.1, color=CODE)
s = content("Imutabilidade corta dos dois lados: desfazer o roubo exigiu bifurcar a rede",
            "diagram_fork.png", size=23, top=1.7, max_h=4.0)
subtitle(s, "a cadeia é imutável → o roubo ficou gravado; a comunidade bifurcou: Ethereum (desfez) e Ethereum Classic (manteve)",
         top=6.1, color=GRAY)

# ---- Etherscan: existe de verdade ----
divider(slide(), "Isto existe de verdade, numa rede pública")
s = content("No Etherscan, o mesmo contrato é público, verificado e permanente",
            "diagram_etherscan.png", size=24, top=1.7, max_h=4.1,
            demo="▶ demo · Etherscan")
subtitle(s, "o escrow da aula, mas movendo bilhões: código verificado, leitura sem carteira, eventos gravados pra sempre",
         top=6.25, color=CHAIN)
s = content("Uma transferência de token: 0 ETH movido, milhares de dólares dentro de uma chamada",
            "diagram_erc20.png", size=22, top=1.7, max_h=4.1,
            demo="▶ demo · Etherscan")
subtitle(s, "Value = 0 ETH, mas a seção de tokens mostra a transferência: o dinheiro andou dentro de transfer()",
         top=6.25, color=GRAY)

# ---- você precisa mesmo de blockchain? ----
divider(slide(), "Você precisa mesmo de blockchain?")
s = content("A árvore de decisão de Wüst–Gervais",
            "fig_decision_tree.png", size=28, top=1.5, max_h=5.5)
s = content("Três casos do dia a dia na mesma árvore",
            "diagram_cases.png", size=27, top=1.75, max_h=4.2)
subtitle(s, "só o caso do meio — sem dono e sem confiança — realmente precisa de blockchain",
         top=6.3, color=BLACK)

# ---- final (síntese das duas aulas) ----
content("Blockchain resolve um problema: concordar sem confiar. Havendo confiança, o simples vence.",
        "fig_synthesis.png", size=22, top=1.7, max_h=4.9)

prs.save(OUTPUT)
print(f"wrote {OUTPUT}  ({len(prs.slides._sldIdLst)} slides)")
